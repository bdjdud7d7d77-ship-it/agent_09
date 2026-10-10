"""Reusable python-pptx helpers for the mk-ppt skill.

Import from a generator script:
    sys.path.insert(0, "<project>/.claude/skills/mk-ppt/scripts")
    from pptx_helpers import *
"""
import os

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt, Emu

__all__ = [
    "FONT", "NAVY", "GRAY", "WHITE", "Inches", "Pt", "RGBColor", "PP_ALIGN",
    "new_deck", "set_font", "add_title_slide", "add_section_slide",
    "add_bullet_slide", "add_two_column_slide", "add_table_slide",
    "add_chart_slide", "add_image_slide", "set_notes", "unique_path", "save",
]

FONT = "맑은 고딕"
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
GRAY = RGBColor(0x59, 0x59, 0x59)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

# Default template layout indices
L_TITLE, L_CONTENT, L_SECTION, L_TITLE_ONLY, L_BLANK = 0, 1, 2, 5, 6

CHART_TYPES = {
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "stacked": XL_CHART_TYPE.COLUMN_STACKED,
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "pie": XL_CHART_TYPE.PIE,
    "doughnut": XL_CHART_TYPE.DOUGHNUT,
    "area": XL_CHART_TYPE.AREA,
    "radar": XL_CHART_TYPE.RADAR,
}


def new_deck(template=None, widescreen=True):
    """New deck (16:9 by default) or one based on a template .pptx."""
    prs = Presentation(template) if template else Presentation()
    if widescreen and not template:
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    return prs


def set_font(run_or_font, size=None, bold=None, color=None, name=FONT):
    """Set font incl. the East Asian slot so Hangul uses `name` too."""
    font = getattr(run_or_font, "font", run_or_font)
    font.name = name
    rpr = font._element if font._element.tag.endswith("rPr") else font._rPr
    ea = rpr.find(qn("a:ea"))
    if ea is None:
        ea = rpr.makeelement(qn("a:ea"), {})
        rpr.append(ea)
    ea.set("typeface", name)
    if size: font.size = Pt(size)
    if bold is not None: font.bold = bold
    if color is not None: font.color.rgb = color


def _style_frame(tf, size, color=None, bold=None):
    for p in tf.paragraphs:
        for r in p.runs:
            set_font(r, size, bold, color)


def _fit_width(prs, shape, margin=Inches(0.6)):
    """Stretch a placeholder across the slide (default layouts are 4:3)."""
    top, height = shape.top, shape.height  # read inherited geometry first
    shape.left, shape.width = margin, prs.slide_width - 2 * margin
    shape.top, shape.height = top, height


def _title(prs, slide, text, size=32):
    t = slide.shapes.title
    _fit_width(prs, t)
    t.text = text
    _style_frame(t.text_frame, size, NAVY, True)
    return t


def _content_box(prs, slide):
    """Area below the title for free-placed content."""
    left, top = Inches(0.6), Inches(1.6)
    return left, top, prs.slide_width - 2 * left, prs.slide_height - top - Inches(0.5)


def set_notes(slide, notes):
    if notes:
        slide.notes_slide.notes_text_frame.text = notes


def add_title_slide(prs, title, subtitle="", notes=None):
    s = prs.slides.add_slide(prs.slide_layouts[L_TITLE])
    _title(prs, s, title, 40)
    sub = s.placeholders[1]
    if subtitle:
        _fit_width(prs, sub, Inches(1.5))
        sub.text = subtitle
        _style_frame(sub.text_frame, 20, GRAY)
    else:
        sub._element.getparent().remove(sub._element)  # no empty placeholder
    set_notes(s, notes)
    return s


def add_section_slide(prs, title, subtitle="", notes=None):
    s = prs.slides.add_slide(prs.slide_layouts[L_SECTION])
    _title(prs, s, title, 36)
    body = s.placeholders[1]
    if subtitle:
        _fit_width(prs, body)
        body.text = subtitle
        _style_frame(body.text_frame, 18, GRAY)
    else:
        body._element.getparent().remove(body._element)
    set_notes(s, notes)
    return s


def _fill_bullets(tf, items, size):
    """items: str or (str, level)."""
    tf.word_wrap = True
    for i, item in enumerate(items):
        text, level = (item, 0) if isinstance(item, str) else item
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text, p.level = text, level
        p.space_after = Pt(6)
        for r in p.runs:
            set_font(r, size - 4 * level)


def add_bullet_slide(prs, title, items, notes=None, size=22):
    s = prs.slides.add_slide(prs.slide_layouts[L_CONTENT])
    _title(prs, s, title)
    body = s.placeholders[1]
    body.left, body.top, body.width, body.height = _content_box(prs, s)
    _fill_bullets(body.text_frame, items, size)
    set_notes(s, notes)
    return s


def add_two_column_slide(prs, title, left_items, right_items,
                         left_head=None, right_head=None, notes=None, size=20):
    s = prs.slides.add_slide(prs.slide_layouts[L_TITLE_ONLY])
    _title(prs, s, title)
    x, y, w, h = _content_box(prs, s)
    gap = Inches(0.4)
    col_w = (w - gap) // 2
    for i, (head, items) in enumerate(((left_head, left_items), (right_head, right_items))):
        box = s.shapes.add_textbox(x + i * (col_w + gap), y, col_w, h)
        tf = box.text_frame
        rows = ([head] if head else []) + list(items)
        _fill_bullets(tf, rows, size)
        if head:
            set_font(tf.paragraphs[0].runs[0], size + 2, True, NAVY)
    set_notes(s, notes)
    return s


def add_table_slide(prs, title, rows, col_widths=None, notes=None, size=14,
                    header_fill=NAVY):
    """rows: list of lists; first row is the header."""
    s = prs.slides.add_slide(prs.slide_layouts[L_TITLE_ONLY])
    _title(prs, s, title)
    x, y, w, _ = _content_box(prs, s)
    n_rows, n_cols = len(rows), len(rows[0])
    row_h = Inches(0.45)
    tbl = s.shapes.add_table(n_rows, n_cols, x, y, w, row_h * n_rows).table
    if col_widths:
        total = sum(col_widths)
        for c, cw in enumerate(col_widths):
            tbl.columns[c].width = Emu(int(w * cw / total))
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.text = str(val)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    set_font(run, size, r == 0, WHITE if r == 0 else None)
            if r == 0 and header_fill is not None:
                cell.fill.solid()
                cell.fill.fore_color.rgb = header_fill
    set_notes(s, notes)
    return s


def add_chart_slide(prs, title, categories, series, kind="column", notes=None,
                    legend=True, labels=False, number_format=None):
    """series: {"name": [values...]}; kind: see CHART_TYPES."""
    s = prs.slides.add_slide(prs.slide_layouts[L_TITLE_ONLY])
    _title(prs, s, title)
    data = CategoryChartData()
    data.categories = categories
    for name, values in series.items():
        data.add_series(name, values)
    x, y, w, h = _content_box(prs, s)
    chart = s.shapes.add_chart(CHART_TYPES[kind], x, y, w, h, data).chart
    chart.font.size = Pt(14)
    chart.font.name = FONT
    chart.has_legend = legend or kind in ("pie", "doughnut")
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
    if labels:
        plot = chart.plots[0]
        plot.has_data_labels = True
        if number_format:
            plot.data_labels.number_format = number_format
            plot.data_labels.number_format_is_linked = False
    set_notes(s, notes)
    return s


def add_image_slide(prs, title, image_path, caption=None, notes=None):
    """Picture scaled to fit the content area, centered."""
    s = prs.slides.add_slide(prs.slide_layouts[L_TITLE_ONLY])
    _title(prs, s, title)
    x, y, w, h = _content_box(prs, s)
    if caption:
        h -= Inches(0.6)
    pic = s.shapes.add_picture(image_path, x, y)
    scale = min(w / pic.width, h / pic.height, 1.0)
    pic.width, pic.height = int(pic.width * scale), int(pic.height * scale)
    pic.left = x + (w - pic.width) // 2
    pic.top = y + (h - pic.height) // 2
    if caption:
        box = s.shapes.add_textbox(x, y + h + Inches(0.1), w, Inches(0.5))
        p = box.text_frame.paragraphs[0]
        p.text, p.alignment = caption, PP_ALIGN.CENTER
        set_font(p.runs[0], 14, color=GRAY)
    set_notes(s, notes)
    return s


def unique_path(out_dir, base, ext=".pptx"):
    """Never overwrite: base.pptx, base_v2.pptx, base_v3.pptx, ..."""
    os.makedirs(out_dir, exist_ok=True)
    path, n = os.path.join(out_dir, base + ext), 2
    while os.path.exists(path):
        path = os.path.join(out_dir, f"{base}_v{n}{ext}")
        n += 1
    return path


def save(prs, out_dir, base):
    path = unique_path(out_dir, base)
    prs.save(path)
    return path
