"""Print a .pptx outline: slides, shapes, text, tables, charts, notes.

Usage: python inspect_pptx.py <file.pptx>
"""
import sys

from pptx import Presentation
from pptx.util import Emu


def inch(v):
    return f"{Emu(v).inches:.2f}in" if v is not None else "?"


def main(path):
    prs = Presentation(path)
    print(f"{path}\nsize: {inch(prs.slide_width)} x {inch(prs.slide_height)}, "
          f"slides: {len(prs.slides)}")
    warnings = []
    for i, slide in enumerate(prs.slides, 1):
        print(f"\n--- slide {i} [{slide.slide_layout.name}] ---")
        for sh in slide.shapes:
            kind = sh.shape_type
            pos = f"@({inch(sh.left)},{inch(sh.top)}) {inch(sh.width)}x{inch(sh.height)}"
            print(f"  [{kind}] {sh.name} {pos}")
            if sh.is_placeholder and sh.has_text_frame and not sh.text_frame.text.strip():
                warnings.append(f"slide {i}: empty placeholder '{sh.name}'")
            if sh.has_text_frame and sh.text_frame.text.strip():
                for p in sh.text_frame.paragraphs:
                    if p.text.strip():
                        print(f"      {'  ' * p.level}- {p.text}")
            if getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    print("      | " + " | ".join(c.text for c in row.cells) + " |")
            if getattr(sh, "has_chart", False) and sh.has_chart:
                ch = sh.chart
                print(f"      chart: {ch.chart_type}, categories={list(ch.plots[0].categories)}")
                for s in ch.series:
                    print(f"        {s.name}: {list(s.values)}")
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            print(f"  notes: {slide.notes_slide.notes_text_frame.text}")
    if warnings:
        print("\nWARNINGS:")
        for w in warnings:
            print("  " + w)


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    sys.stdout.reconfigure(encoding="utf-8")
    main(sys.argv[1])
