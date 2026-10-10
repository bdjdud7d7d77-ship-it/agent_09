"""Generate a 5-slide Korean deck introducing Anthropic, the company behind Claude.

Main color: orange. Run from any directory:
    python scripts/make_anthropic_company_ppt.py
"""
import os
import sys

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
sys.path.insert(0, os.path.join(ROOT, ".claude", "skills", "mk-ppt", "scripts"))

import pptx_helpers
from pptx_helpers import *  # noqa: F401,F403
from pptx.enum.shapes import MSO_SHAPE

ORANGE = RGBColor(0xE8, 0x77, 0x2E)
DARK = RGBColor(0x33, 0x33, 0x33)
LIGHT = RGBColor(0xFD, 0xF1, 0xE8)

# Helpers color titles with the module-level NAVY; switch it to the main color.
pptx_helpers.NAVY = ORANGE


def add_rect(slide, left, top, width, height, color, back=True):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    if back:  # send to back so it never covers text
        tree = slide.shapes._spTree
        tree.remove(shape._element)
        tree.insert(2, shape._element)
    return shape


def accent(prs, slide, page):
    """Orange top bar, thin bottom line and page number."""
    add_rect(slide, 0, 0, prs.slide_width, Inches(0.18), ORANGE)
    add_rect(slide, Inches(0.6), prs.slide_height - Inches(0.45),
             prs.slide_width - Inches(1.2), Pt(1.5), ORANGE)
    box = slide.shapes.add_textbox(prs.slide_width - Inches(1.6), prs.slide_height - Inches(0.42),
                                   Inches(1.0), Inches(0.35))
    p = box.text_frame.paragraphs[0]
    p.text, p.alignment = f"{page} / 5", PP_ALIGN.RIGHT
    set_font(p.runs[0], 12, color=GRAY)


def color_bullets(slide):
    """Level-0 bullets dark, level-1 gray."""
    for p in slide.placeholders[1].text_frame.paragraphs:
        for r in p.runs:
            r.font.color.rgb = DARK if p.level == 0 else GRAY


prs = new_deck()

# 1. Title
s = add_title_slide(
    prs, "Anthropic", "Claude를 만든 AI 안전 연구 기업",
    notes="오늘은 Claude를 개발한 회사 Anthropic을 소개합니다. "
          "회사 개요, 핵심 제품, 안전 철학 순서로 살펴보겠습니다.")
add_rect(s, 0, 0, prs.slide_width, prs.slide_height, ORANGE)
add_rect(s, Inches(1.5), Inches(4.05), prs.slide_width - Inches(3.0), Pt(3), WHITE, back=False)
for shape in (s.shapes.title, s.placeholders[1]):
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            r.font.color.rgb = WHITE
set_font(s.shapes.title.text_frame.paragraphs[0].runs[0], 54, True, WHITE)

# 2. Company overview
s = add_bullet_slide(prs, "회사 개요", [
    "2021년 설립된 AI 안전·연구 기업",
    ("미국 샌프란시스코 본사", 1),
    "공동 창업자: 다리오 아모데이(CEO), 다니엘라 아모데이(사장) 등",
    ("OpenAI 출신 연구자들이 중심이 되어 창업", 1),
    "공익 법인(Public Benefit Corporation) 형태로 운영",
    "미션: 안전하고 신뢰할 수 있으며 유익한 AI 구축",
], notes="Anthropic은 2021년 설립되었고, 이익뿐 아니라 공익을 함께 추구하는 "
         "공익 법인 형태를 택했습니다. 미션의 핵심은 '안전한 AI'입니다.")
color_bullets(s)
accent(prs, s, 2)

# 3. Core product: Claude
s = add_table_slide(prs, "핵심 제품: Claude", [
    ["구분", "내용"],
    ["모델 제품군", "Opus(최고 성능) · Sonnet(성능·속도 균형) · Haiku(빠르고 경제적)"],
    ["Claude 앱", "웹 · 데스크톱 · 모바일에서 대화형 AI 비서로 사용"],
    ["Claude Code", "터미널·IDE에서 코드 작성과 수정을 돕는 에이전트형 코딩 도구"],
    ["Claude API", "개발자가 자체 서비스에 Claude를 연동 (AWS Bedrock, Google Vertex AI로도 제공)"],
], col_widths=[1, 3.2], size=18, header_fill=ORANGE,
    notes="Claude는 용도에 따라 Opus, Sonnet, Haiku 세 가지 등급으로 제공됩니다. "
          "일반 사용자는 앱으로, 개발자는 API와 Claude Code로 활용합니다.")
tbl = next(sh for sh in s.shapes if sh.has_table).table
for r in range(1, len(tbl.rows)):
    tbl.rows[r].height = Inches(0.8)
    for c in range(len(tbl.columns)):
        cell = tbl.cell(r, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = LIGHT if r % 2 else WHITE
        for p in cell.text_frame.paragraphs:
            for run in p.runs:
                run.font.color.rgb = DARK
                run.font.bold = c == 0
accent(prs, s, 3)

# 4. Safety philosophy
s = add_two_column_slide(
    prs, "안전 철학과 연구",
    ["원칙(헌법)에 따라 AI가 스스로 답변을 점검·개선",
     "사람의 가치를 명시적인 규칙으로 반영",
     "모델 능력 수준에 맞춰 안전 조치를 단계적으로 강화(RSP)"],
    ["AI 내부 작동 원리를 해석하는 연구(해석 가능성)",
     "연구 성과와 정책 제안을 외부에 공개",
     "장기 공익 신탁(LTBT)으로 지배구조의 안전장치 마련"],
    left_head="Constitutional AI · RSP", right_head="연구와 거버넌스",
    notes="Anthropic의 차별점은 안전입니다. Constitutional AI로 모델의 행동 원칙을 정하고, "
          "책임 있는 확장 정책(RSP)으로 능력이 커질수록 안전 기준을 높입니다.")
for sh in s.shapes:
    if sh.has_text_frame and sh != s.shapes.title:
        for i, p in enumerate(sh.text_frame.paragraphs):
            if i:
                for r in p.runs:
                    r.font.color.rgb = DARK
accent(prs, s, 4)

# 5. Summary
s = add_bullet_slide(prs, "정리", [
    "Anthropic = '안전'을 최우선으로 하는 AI 연구 기업",
    "Claude: 앱 · Claude Code · API로 개인과 기업 모두 활용",
    "Constitutional AI와 RSP로 신뢰할 수 있는 AI를 추구",
], notes="세 가지만 기억해 주세요. 안전 중심의 회사, 다양한 형태의 Claude, "
         "그리고 신뢰를 위한 연구입니다. 감사합니다.", size=26)
color_bullets(s)
box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.17), Inches(5.3),
                         Inches(5.0), Inches(0.9))
box.fill.solid()
box.fill.fore_color.rgb = ORANGE
box.line.fill.background()
p = box.text_frame.paragraphs[0]
p.text, p.alignment = "감사합니다", PP_ALIGN.CENTER
set_font(p.runs[0], 28, True, WHITE)
accent(prs, s, 5)

out = save(prs, os.path.join(ROOT, "reports", "anthropic-company"), "anthropic_company_slides")
print(out)
