"""Generate reports/ai-development/ai_development_report.docx from the research notes."""
import os
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "reports", "ai-development")
os.makedirs(OUT_DIR, exist_ok=True)

# No overwrite: pick a new filename if one already exists
base = "ai_development_report"
path = os.path.join(OUT_DIR, base + ".docx")
n = 2
while os.path.exists(path):
    path = os.path.join(OUT_DIR, f"{base}_v{n}.docx")
    n += 1

FONT = "맑은 고딕"
NAVY = RGBColor(0x1F, 0x3A, 0x5F)

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.5)
sec.top_margin = sec.bottom_margin = Cm(2.5)


def set_font(run, size=None, bold=None, color=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    if size: run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if color: run.font.color.rgb = color


style = doc.styles["Normal"]
style.font.name = FONT
style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
style.font.size = Pt(11)
style.paragraph_format.line_spacing = 1.5
style.paragraph_format.space_after = Pt(6)
for lvl, size in ((1, 16), (2, 13)):
    hs = doc.styles[f"Heading {lvl}"]
    hs.font.name = FONT
    hs.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    hs.font.size = Pt(size)
    hs.font.color.rgb = NAVY
    hs.font.bold = True


def h(text, lvl):
    doc.add_heading(text, level=lvl)


def p(text):
    para = doc.add_paragraph()
    para.paragraph_format.first_line_indent = Cm(0.5)
    para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_font(para.add_run(text))


def bullets(items):
    for it in items:
        para = doc.add_paragraph(style="List Bullet")
        if isinstance(it, tuple):
            set_font(para.add_run(it[0]), bold=True)
            set_font(para.add_run(it[1]))
        else:
            set_font(para.add_run(it))


def shade(cell, hex_color):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def table(header, rows, widths=None, caption=None):
    if caption:
        cp = doc.add_paragraph()
        set_font(cp.add_run(caption), size=10, bold=True)
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, txt in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        set_font(c.paragraphs[0].add_run(txt), size=10, bold=True, color=RGBColor(255, 255, 255))
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        shade(c, "1F3A5F")
    for r_i, row in enumerate(rows):
        cells = t.add_row().cells
        for i, txt in enumerate(row):
            cells[i].text = ""
            cells[i].paragraphs[0].paragraph_format.line_spacing = 1.2
            set_font(cells[i].paragraphs[0].add_run(txt), size=9.5)
            if r_i % 2 == 1:
                shade(cells[i], "EEF2F7")
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()


# ---------- Cover ----------
for _ in range(6):
    doc.add_paragraph()
tp = doc.add_paragraph(); tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(tp.add_run("AI 발전 보고서"), size=28, bold=True, color=NAVY)
sp = doc.add_paragraph(); sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(sp.add_run("인공지능의 역사, 최신 동향, 사회적 영향과 미래 전망"), size=14, color=RGBColor(0x55, 0x55, 0x55))
for _ in range(10):
    doc.add_paragraph()
dp = doc.add_paragraph(); dp.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(dp.add_run("작성일: 2026년 10월 4일"), size=12)
doc.add_page_break()

# ---------- Table of contents ----------
h("목차", 1)
for line in ["Ⅰ. 서론", "    1. 작성 배경 및 목적", "    2. 인공지능의 정의",
             "Ⅱ. 본론", "    1. AI 발전의 역사", "    2. 최신 기술 동향", "    3. 산업별 활용 사례",
             "    4. 사회·경제적 영향과 윤리·규제", "Ⅲ. 결론", "    1. 요약", "    2. 향후 전망", "    3. 시사점", "참고문헌"]:
    para = doc.add_paragraph(); para.paragraph_format.space_after = Pt(2)
    set_font(para.add_run(line), bold=not line.startswith(" "))
doc.add_page_break()

# ---------- I. Introduction ----------
h("Ⅰ. 서론", 1)
h("1. 작성 배경 및 목적", 2)
p("2022년 11월 ChatGPT가 공개된 뒤 인공지능(AI)은 일부 연구자와 기업의 기술을 넘어 일상과 산업 전반에 빠르게 퍼졌다. "
  "스탠퍼드 대학교 인간중심AI연구소(HAI)의 「2026 AI 인덱스 보고서」에 따르면 생성형 AI는 출시 3년 만에 세계 인구의 53%가 쓰게 되어 "
  "개인용 컴퓨터나 인터넷보다 빠르게 확산되었고, 전 세계 조직의 88%가 어떤 형태로든 AI를 도입했다.")
p("이처럼 AI는 산업 구조, 일자리, 교육, 정책에 큰 변화를 가져오고 있다. 그러나 기술 발전 속도에 비해 안전, 규제, 사회적 대응은 뒤처지고 있다는 지적도 많다. "
  "이 보고서는 AI 발전의 역사와 최신 기술 동향을 정리하고, 산업별 활용 사례와 사회·경제적 영향을 살펴본 뒤, 앞으로의 전망과 시사점을 제시하는 것을 목적으로 한다.")
h("2. 인공지능의 정의", 2)
p("인공지능은 학습, 추론, 지각, 언어 이해 등 사람의 지적 능력을 컴퓨터로 구현하는 기술을 말한다. "
  "이 용어는 1956년 미국 다트머스 대학교에서 열린 워크숍에서 존 매카시가 처음 사용했다. 당시 연구자들은 "
  "\"학습을 비롯한 지능의 모든 특성은 원칙적으로 기계가 흉내 낼 수 있을 만큼 정확하게 기술할 수 있다\"고 보았다. "
  "오늘날 AI는 데이터에서 스스로 패턴을 배우는 머신러닝과, 이를 다층 신경망으로 확장한 딥러닝을 중심으로 발전하고 있다.")

# ---------- II. Body ----------
h("Ⅱ. 본론", 1)
h("1. AI 발전의 역사", 2)
p("AI의 역사는 기대와 투자가 몰리는 '여름'과 실망과 자금 삭감이 이어지는 '겨울'이 반복된 과정이었다. "
  "초기에는 사람이 규칙을 직접 입력하는 기호주의 AI가 주류였으나, 컴퓨팅 성능의 한계와 지나친 기대 때문에 1974–1980년과 1987–1993년 두 차례 'AI 겨울'을 겪었다.")
p("전환점은 2012년이었다. 제프리 힌턴 연구팀의 AlexNet이 이미지 인식 대회(ImageNet)에서 2위와 약 10.8%p 차이로 우승하면서 딥러닝 시대가 열렸다. "
  "2016년에는 알파고가 이세돌 9단을 이겼고, 2017년 발표된 트랜스포머 구조는 오늘날 대규모 언어 모델(LLM)의 기반이 되었다. "
  "2022년 ChatGPT는 두 달 만에 사용자 1억 명을 모아 생성형 AI 대중화를 이끌었다. "
  "2024년에는 힌턴이 노벨 물리학상을, 알파폴드를 만든 허사비스와 점퍼가 노벨 화학상을 받아 AI의 학문적 기여가 공식적으로 인정받았다.")
table(["시기", "주요 사건", "의의"], [
    ["1956", "다트머스 워크숍", "'인공지능' 용어 탄생, 학문 분야 출범"],
    ["1974–1993", "두 차례의 AI 겨울", "과도한 기대와 기술 한계로 연구 침체"],
    ["1997", "IBM 딥블루, 체스 챔피언에게 승리", "특정 분야에서 인간을 넘어선 AI"],
    ["2012", "AlexNet, ImageNet 우승", "딥러닝 붐의 시작"],
    ["2016", "알파고, 이세돌에게 승리", "복잡한 직관의 영역으로 여겨진 바둑 정복"],
    ["2017", "트랜스포머 발표", "현대 대규모 언어 모델의 기반"],
    ["2022", "ChatGPT 공개", "생성형 AI 대중화 (2개월 만에 사용자 1억 명)"],
    ["2024", "AI 연구에 노벨 물리학상·화학상", "AI의 과학적 기여 공식 인정"],
    ["2025–2026", "추론 모델과 AI 에이전트 확산", "답변하는 AI에서 일하는 AI로"],
], widths=[2.5, 6, 7.5], caption="[표 1] AI 발전의 주요 이정표")

h("2. 최신 기술 동향", 2)
p("「2026 AI 인덱스 보고서」는 AI 성능이 정체되지 않고 오히려 빨라지고 있다고 평가했다. "
  "대표적인 코딩 평가인 SWE-bench Verified의 최고 점수는 1년 만에 약 60%에서 100% 가까이로 올랐고, "
  "여러 모델이 박사 수준 과학 문제와 수학 경시대회에서 인간 기준치를 넘어섰다. 최근 기술 동향은 다음 네 가지로 정리할 수 있다.")
bullets([
    ("추론 모델: ", "OpenAI o1, DeepSeek-R1 등은 답하기 전에 여러 단계로 '생각'하는 과정에 연산을 더 써서 수학·과학·코딩 정확도를 높인다."),
    ("멀티모달 AI: ", "하나의 모델이 텍스트, 이미지, 음성, 영상을 함께 이해하고 생성한다. 이 기능은 주요 클라우드 서비스에 기본으로 들어가 있다."),
    ("AI 에이전트: ", "질문에 답하는 도구에서 계획을 세우고 도구를 쓰며 일을 직접 처리하는 에이전트로 바뀌고 있다. 전문 에이전트 여러 개가 협업하는 멀티 에이전트 시스템도 늘고 있다."),
    ("미·중 기술 격차 해소: ", "2025년 초부터 미국과 중국 모델이 선두를 주고받았고, 2026년 3월 기준 격차는 2.7%에 불과하다."),
])
p("한편 AI의 능력은 고르지 않다. 국제수학올림피아드에서 금메달 수준을 낸 모델이 아날로그 시계는 약 50%만 맞게 읽는다. "
  "이를 '들쭉날쭉한 최전선(Jagged Frontier)'이라 부르며, 벤치마크 점수만으로 실제 업무 성능을 판단하기 어렵다는 점을 보여 준다.")

h("3. 산업별 활용 사례", 2)
p("AI는 거의 모든 산업으로 퍼지고 있다. 다만 산업별 도입률은 조사 기관마다 차이가 있어, 아래 수치는 대략적인 경향으로 보는 것이 바람직하다.")
table(["산업", "주요 활용", "현황(참고 수치)"], [
    ["의료", "의료영상 분석, 임상 의사결정 지원, 신약 개발", "도입률 약 62–67%, 미국 FDA AI 의료기기 승인 2025년 약 178건"],
    ["금융", "사기 탐지, 신용평가, 리스크 관리, 고객 상담", "도입률 약 79%, 기업 AI 지출 비중 1위"],
    ["제조", "예지 정비, 품질 검사, 공정 최적화", "제조기업 약 77% 활용, AI 지출 전년 대비 약 48% 증가"],
    ["교육", "맞춤형 학습, 채점 지원, 콘텐츠 제작", "도입률 약 34%로 최저, 대학생 5명 중 4명은 생성형 AI 사용"],
    ["소프트웨어", "코드 생성, 테스트, 자율 코딩 에이전트", "기술 분야 도입률 약 88%"],
], widths=[2.5, 6, 7.5], caption="[표 2] 산업별 AI 활용 현황")
p("과학 연구에서도 AI의 역할이 커지고 있다. 딥마인드의 알파폴드는 50년 넘게 풀리지 않던 단백질 구조 예측 문제를 해결하고 "
  "알려진 단백질 약 2억 개의 구조를 예측했다. 이는 신약 개발과 생명과학 연구 속도를 크게 높이고 있다. "
  "다만 도입률이 높다고 성과가 바로 나는 것은 아니어서, 상당수 기업이 아직 뚜렷한 투자 수익을 얻지 못하고 있다는 조사도 있다.")

h("4. 사회·경제적 영향과 윤리·규제", 2)
p("경제적 측면에서 AI는 막대한 자본을 끌어들이고 있다. 전 세계 기업의 AI 투자는 5,817억 달러로 전년보다 130% 늘었고, "
  "그중 생성형 AI 투자는 1,709억 달러로 약 5배 증가했다. 특히 미국의 민간 AI 투자(2,859억 달러)는 중국의 23배가 넘는다.")
p("일자리 측면에서는 기회와 위협이 함께 나타난다. 세계경제포럼(WEF)은 2030년까지 일자리 1억 7천만 개가 생기고 9,200만 개가 사라져 "
  "순증가는 7,800만 개가 될 것으로 전망했다. 그러나 전 세계 노동자의 59%가 재교육이 필요하다. "
  "국제통화기금(IMF)은 전 세계 일자리의 약 40%, 선진국은 약 60%가 AI의 영향을 받는다고 분석했다. "
  "AI가 일자리에 긍정적일 것이라는 응답은 전문가 73%, 일반 대중 23%로 인식 차이도 크다.")
p("윤리와 안전 문제도 커지고 있다. 기록된 AI 사고는 2024년 233건에서 362건으로 늘었고, 딥페이크와 허위정보, 편향, 개인정보 침해, "
  "학습 데이터 저작권, 환각(그럴듯한 오답) 등이 주요 문제로 꼽힌다. 에너지 소비도 과제다. 국제에너지기구(IEA)는 데이터센터 전력 사용량이 "
  "2024년 약 415TWh에서 2030년 약 945TWh로 두 배 이상 늘고, AI 전용 서버의 전력 사용량은 약 3배가 될 것으로 내다봤다.")
p("이에 각국은 AI 규제를 본격화하고 있다. EU는 AI법(AI Act)을 위험 수준에 따라 단계적으로 시행하고 있으며, 대부분의 조항이 2026년 8월부터 적용된다. "
  "한국은 2026년 1월 22일 「인공지능 발전과 신뢰 기반 조성 등에 관한 기본법(AI 기본법)」을 시행해 EU에 이어 세계에서 두 번째로 포괄적인 AI 법을 갖추었다. "
  "이 법은 고영향 AI와 생성형 AI에 위험 평가 의무를 두면서도, 최소 1년의 과태료 유예 기간을 두어 혁신과 안전의 균형을 꾀하고 있다.")
table(["구분", "EU AI법", "한국 AI 기본법"], [
    ["시행 시기", "2025년 8월부터 단계적 (대부분 2026년 8월, 고위험 AI 2027년 8월)", "2026년 1월 22일"],
    ["접근 방식", "위험 수준별 규제", "산업 진흥과 신뢰 확보 병행"],
    ["주요 내용", "범용 AI 의무, 고위험 AI 요건, 금지 AI 지정", "고영향·생성형 AI 위험 평가, 해외 기업 국내 대리인 지정, 과태료 최소 1년 유예"],
], widths=[3, 6.5, 6.5], caption="[표 3] EU와 한국의 AI 규제 비교")

# ---------- III. Conclusion ----------
h("Ⅲ. 결론", 1)
h("1. 요약", 2)
p("AI는 1956년 학문 분야로 출발한 뒤 두 차례의 침체기를 거쳐, 딥러닝과 트랜스포머, 생성형 AI를 계기로 PC나 인터넷보다 빠르게 퍼진 범용 기술이 되었다. "
  "현재는 추론 모델, 멀티모달 AI, AI 에이전트를 중심으로 성능이 빠르게 오르고 있으며, 의료·금융·제조·소프트웨어 등 거의 모든 산업에서 쓰이고 있다. "
  "반면 안전사고 증가, 일자리 변화, 에너지 소비, 규제 공백 같은 문제도 함께 커지고 있다.")
h("2. 향후 전망", 2)
bullets([
    ("에이전트 중심으로의 전환: ", "AI는 답하는 도구에서 일을 수행하는 '디지털 동료'로 바뀌고, 사람은 목표 설정과 감독, 검증을 맡게 될 것이다."),
    ("인프라와 에너지 제약: ", "전력, 반도체, 전력망 용량이 AI 성장의 주요 한계가 되며, 에너지 효율과 청정 전력 확보가 경쟁력이 될 것이다."),
    ("규제의 본격화: ", "2026년을 기점으로 한국과 EU의 AI 법이 시행되면서, 기업은 AI 안전과 투명성 확보를 필수 과제로 다루게 될 것이다."),
    ("노동시장 재편: ", "일자리는 순증가할 전망이지만, 대규모 직무 전환과 재교육이 불가피하다."),
])
h("3. 시사점", 2)
p("AI 발전의 핵심 과제는 빠르게 앞서가는 기술과 뒤처진 사회적 대응 사이의 격차를 줄이는 것이다. "
  "정부는 AI 기본법을 실효성 있게 운영하면서 산업 경쟁력을 키울 수 있도록 제도를 다듬어야 한다. "
  "기업은 단순 도입을 넘어 실제 성과로 이어지도록 업무 방식을 바꾸고, 안전과 윤리 기준을 갖춰야 한다. "
  "교육 현장은 AI 리터러시 교육을 강화하고, 개인은 AI를 잘 활용하는 능력과 결과를 비판적으로 검증하는 능력을 함께 길러야 한다. "
  "AI는 이미 거스를 수 없는 흐름이며, 이를 사회 전체의 이익으로 연결하려는 노력이 그 어느 때보다 중요하다.")

# ---------- References ----------
h("참고문헌", 1)
refs = [
    "Stanford HAI, \"The 2026 AI Index Report\", https://hai.stanford.edu/ai-index/2026-ai-index-report",
    "Wikipedia, \"History of artificial intelligence\", https://en.wikipedia.org/wiki/History_of_artificial_intelligence",
    "NobelPrize.org, \"The Nobel Prize in Chemistry 2024 – Popular information\", https://www.nobelprize.org/prizes/chemistry/2024/popular-information/",
    "University of Toronto, \"Geoffrey Hinton wins Nobel Prize in Physics\", https://www.utoronto.ca/news/geoffrey-hinton-wins-nobel-prize",
    "Nature, \"Chemistry Nobel goes to developers of AlphaFold AI\", https://www.nature.com/articles/d41586-024-03214-7",
    "World Economic Forum, \"The future of jobs: 6 decision-makers on AI and talent strategies\", https://www.weforum.org/stories/2026/01/how-ai-will-affect-work-in-different-industries/",
    "IMF, \"Bridging Skill Gaps for the Future: New Jobs Creation in the AI Age\" (2026), https://www.imf.org/-/media/files/publications/sdn/2026/english/sdnea2026001.pdf",
    "DCD, \"IEA: Data center energy consumption set to double by 2030 to 945TWh\", https://www.datacenterdynamics.com/en/news/iea-data-center-energy-consumption-set-to-double-by-2030-to-945twh/",
    "Stimson Center, \"South Korea's AI Basic Act\", https://www.stimson.org/2026/south-koreas-ai-basic-act-seeking-balance-between-industry-innovation-and-social-risk/",
    "Cooley, \"South Korea's AI Basic Act: Overview and Key Takeaways\", https://www.cooley.com/news/insight/2026/2026-01-27-south-koreas-ai-basic-act-overview-and-key-takeaways",
    "Hugging Face, \"AI Trends 2026: Test-Time Reasoning and the Rise of Reflective Agents\", https://huggingface.co/blog/aufklarer/ai-trends-2026-test-time-reasoning-reflective-agen",
    "Axis Intelligence, \"AI Adoption by Industry 2026\", https://axis-intelligence.com/ai-adoption-by-industry-statistics/",
]
for i, r in enumerate(refs, 1):
    para = doc.add_paragraph()
    para.paragraph_format.line_spacing = 1.2
    set_font(para.add_run(f"[{i}] {r}"), size=9.5)

# Page numbers in footer
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = fp.add_run()
for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
    if tag:
        el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), tag)
    else:
        el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = text
    run._r.append(el)

doc.save(path)
print("saved:", path)
